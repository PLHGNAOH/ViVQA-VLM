"""
make_audit_sheet.py — Tạo bảng duyệt tay (Excel) + thư mục ảnh cho W07: audit dữ liệu và phân tích lỗi.

Vấn đề: đề cương bắt buộc error analysis + hallucination analysis + thách thức tiếng Việt (dấu thanh,
code-switching, OCR), và GVHD yêu cầu "hiểu dataset + nhiễu". Các con số tự động (EM, containment)
không cho biết LỖI THUỘC LOẠI GÌ — phải có người xem ảnh. Script này chuẩn bị đúng 400 dòng để duyệt,
chọn tất định (seed 42), kèm ảnh, có ô chọn nhãn sẵn (dropdown) để nhiều người duyệt song song.

  Sheet "data_audit"  (200 dòng): mẫu TRAIN — nhãn gốc có đúng không? (ViVQA: 25/loại câu hỏi × 4; ViTextVQA: 100)
  Sheet "error_audit" (200 dòng): câu TEST mà B2 trả lời sai EM — lỗi thuộc loại nào? (100 ViVQA + 100 ViTextVQA)
  Sheet "huong_dan": định nghĩa nhãn + phân công gợi ý.
  Ảnh: <out_dir>/images/<dataset>/<tên file> (ViVQA chép từ Drive; ViTextVQA rút từng file từ zip, không giải nén cả 5,6 GB).

Chạy (máy local có Google Drive for desktop, từ gốc repo):
    python -m src.data.make_audit_sheet --drive_root "G:/My Drive/ViVQA-VLM" --out_dir "G:/My Drive/ViVQA-VLM/audit/W07"
Kiểm thử: python -m src.data.make_audit_sheet --selftest
"""
from __future__ import annotations

import argparse
import json
import os
import random
import shutil
import sys
import zipfile
from typing import Any, Dict, List, Optional

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)
from src.eval.metrics import anls, exact_match  # noqa: E402

SEED = 42
RUNS = {
    "vivqa": {"B1": "W07_B1_zeroshot_qwen25vl_vivqa_test_bf16_L4_seed42",
              "B2": "W07_B2_eval_qlora_r16_ep1_vivqa_test_bf16_L4_seed42"},
    "vitextvqa": {"B1": "W07_B1_zeroshot_qwen25vl_vitextvqa_test2000_bf16_L4_seed42",
                  "B2": "W07_B2_eval_qlora_r16_ep1_vitextvqa_test2000_bf16_L4_seed42",
                  "R1": "W07_R1_vitextvqa_test2000_qwen2vl7b_paper_bf16_L4_seed42"},
}
DATA = {"vivqa": {"train": "data/vivqa/train.json", "test": "data/vivqa/test.json", "images": "data/vivqa/images"},
        "vitextvqa": {"train": "data/vitextvqa_official/train.json",
                      "test": "data/vitextvqa_official/test_subset2000_seed42.json",
                      "zip": "data/vitextvqa_official/ViTextVQA_images.zip"}}
DATA_LABELS = ["OK", "Dịch sai / dịch máy (vd 'tía')", "Sai dấu thanh / chính tả", "Đáp án không khớp ảnh",
               "Câu hỏi mơ hồ / nhiều đáp án đúng", "Đáp án không chuẩn hoá (dài, thừa chữ)", "Khác"]
ERROR_LABELS = ["Định dạng / độ dài (ý đúng)", "Đồng nghĩa / khác từ (ý đúng)", "Đọc chữ sai (OCR)",
                "Sai dấu thanh", "Sai vùng / không có trong ảnh (hallucination)", "Code-switching (tiếng Anh)",
                "Đếm sai", "Gold sai / nhiễu", "Khác"]
GUIDE = [
    ("Mục đích", "Duyệt tay 400 mẫu để (1) đo nhiễu nhãn của dataset, (2) phân loại lỗi của B2 cho error/hallucination analysis."),
    ("Cách làm", "Mở ảnh bằng link ở cột 'anh' (hoặc tìm theo tên file trong thư mục images/). Chọn nhãn ở các cột màu vàng bằng dropdown; ghi chú nếu cần."),
    ("data_audit", "Mẫu TRAIN. 'nhan_goc' = nhãn gốc đúng/sai/mơ hồ so với ảnh; 'loai_van_de' theo danh sách; nếu sai ghi đáp án đúng."),
    ("error_audit", "Câu TEST B2 sai EM. 'loai_loi' = nguyên nhân chính (1 nhãn); 'b2_dung_nghia' = Có nếu người đọc chấp nhận câu trả lời của B2; 'hallucination' = Có nếu B2 nói thứ KHÔNG có trong ảnh."),
    ("Nhãn data_audit", " | ".join(DATA_LABELS)),
    ("Nhãn error_audit", " | ".join(ERROR_LABELS)),
    ("Phân công gợi ý", "SV2 (Dataset/OCR): data_audit. SV4 (Evaluation) + Hoàng: error_audit. ≈30 giây/dòng -> ≈1,5–2 giờ/200 dòng."),
    ("Quy tắc", "Không sửa các cột không tô vàng. Mỗi người điền tên ở cột 'nguoi_duyet'. Dòng khó: chọn 'Khác' + ghi chú, không đoán."),
    ("Chọn mẫu", f"Tất định, seed {SEED}: data_audit ViVQA 25 mẫu/loại câu hỏi (color, location, number, object), ViTextVQA 100 mẫu ngẫu nhiên; error_audit 100 câu B2 sai EM mỗi dataset."),
]


def _load(path: str) -> List[Dict[str, Any]]:
    raw = json.load(open(path, encoding="utf-8"))
    if isinstance(raw, dict):
        raw = raw["data"] if isinstance(raw.get("data"), list) else list(raw.values())
    return raw


def _answers(r: Dict[str, Any]) -> List[str]:
    a = r.get("answers", r.get("answer", ""))
    a = a if isinstance(a, list) else [a]
    return [str(x.get("answer", "") if isinstance(x, dict) else x) for x in a if str(x).strip()]


def pick_train(train: List[Dict[str, Any]], dataset: str, seed: int = SEED) -> List[Dict[str, Any]]:
    rng = random.Random(seed)
    if dataset == "vivqa":
        out = []
        for t in ("color", "location", "number", "object"):
            g = [r for r in train if r.get("question_type") == t]
            out += rng.sample(g, min(25, len(g)))
        return out
    return rng.sample(train, min(100, len(train)))


def pick_errors(test: List[Dict[str, Any]], preds: Dict[str, List[Dict[str, Any]]], n: int = 100,
                seed: int = SEED) -> List[Dict[str, Any]]:
    """Câu test B2 sai EM (chọn ngẫu nhiên n câu, seed cố định), kèm dự đoán các run."""
    by = {k: {p["question_id"]: p for p in v} for k, v in preds.items()}
    test_by = {str(r["question_id"]): r for r in test}
    wrong = [q for q, p in by["B2"].items() if not exact_match(p["prediction"], p["answers"])]
    rows = []
    for q in random.Random(seed).sample(sorted(wrong), min(n, len(wrong))):
        t = test_by[q]
        row = {"question_id": q, "image": os.path.basename(str(t["image"])), "question": t["question"],
               "gold": " || ".join(by["B2"][q]["answers"]),
               "question_type": t.get("question_type", "unknown"),
               "ANLS_B2": round(anls(by["B2"][q]["prediction"], by["B2"][q]["answers"]), 3)}
        for k in preds:
            row[k] = by[k][q]["prediction"]
        rows.append(row)
    return rows


def copy_images(names: List[str], dataset: str, drive_root: str, out_dir: str) -> Dict[str, str]:
    """Chép ảnh cần duyệt vào <out_dir>/images/<dataset>/. Trả {tên: đường dẫn tương đối} (thiếu -> '')."""
    dst = os.path.join(out_dir, "images", dataset)
    os.makedirs(dst, exist_ok=True)
    want, got = sorted(set(names)), {}
    if dataset == "vivqa":
        src = os.path.join(drive_root, DATA["vivqa"]["images"])
        for n in want:
            if os.path.exists(os.path.join(src, n)):
                if not os.path.exists(os.path.join(dst, n)):
                    shutil.copyfile(os.path.join(src, n), os.path.join(dst, n))
                got[n] = f"images/{dataset}/{n}"
    else:
        with zipfile.ZipFile(os.path.join(drive_root, DATA["vitextvqa"]["zip"])) as zf:
            members = {os.path.basename(m): m for m in zf.namelist() if not m.endswith("/")}
            for n in want:
                if n in members:
                    if not os.path.exists(os.path.join(dst, n)):
                        with zf.open(members[n]) as fi, open(os.path.join(dst, n), "wb") as fo:
                            shutil.copyfileobj(fi, fo)
                    got[n] = f"images/{dataset}/{n}"
    return {n: got.get(n, "") for n in want}


def write_xlsx(path: str, data_rows: List[Dict[str, Any]], err_rows: List[Dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    yellow = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()
    g = wb.active
    g.title = "huong_dan"
    for i, (k, v) in enumerate(GUIDE, 1):
        g.cell(i, 1, k).font = Font(bold=True)
        g.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    g.column_dimensions["A"].width, g.column_dimensions["B"].width = 20, 120

    def sheet(name, cols, human, rows, lists):
        ws = wb.create_sheet(name)
        ws.append(cols + human)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in cols] + [""] * len(human))
            link = r.get("_link")
            if link:
                cell = ws.cell(ws.max_row, cols.index("anh") + 1)
                cell.hyperlink, cell.style = link, "Hyperlink"
        n = ws.max_row
        for j, h in enumerate(human, len(cols) + 1):
            col = ws.cell(1, j).column_letter
            for i in range(2, n + 1):
                ws.cell(i, j).fill = yellow
            if h in lists:
                dv = DataValidation(type="list", formula1='"' + ",".join(lists[h]) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{n}")
        for j, c in enumerate(cols + human, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = 45 if c in ("question", "gold", "answers", "B1", "B2", "R1") else 16
        ws.freeze_panes = "A2"

    sheet("data_audit", ["stt", "dataset", "question_id", "question_type", "anh", "question", "answers"],
          ["nhan_goc", "loai_van_de", "dap_an_dung_neu_sai", "ghi_chu", "nguoi_duyet"], data_rows,
          {"nhan_goc": ["Đúng", "Sai", "Mơ hồ"], "loai_van_de": [x.replace(",", ";") for x in DATA_LABELS]})
    sheet("error_audit", ["stt", "dataset", "question_id", "question_type", "anh", "question", "gold", "B1", "B2", "R1", "ANLS_B2"],
          ["loai_loi", "b2_dung_nghia", "hallucination", "ghi_chu", "nguoi_duyet"], err_rows,
          {"loai_loi": [x.replace(",", ";") for x in ERROR_LABELS], "b2_dung_nghia": ["Có", "Không"],
           "hallucination": ["Có", "Không"]})
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def build(drive_root: str, out_dir: str, exp_root: str) -> Dict[str, Any]:
    data_rows, err_rows, missing = [], [], []
    for ds in ("vivqa", "vitextvqa"):
        train = _load(os.path.join(drive_root, DATA[ds]["train"]))
        test = _load(os.path.join(drive_root, DATA[ds]["test"]))
        preds = {k: json.load(open(os.path.join(exp_root, r, "predictions.json"), encoding="utf-8"))
                 for k, r in RUNS[ds].items()}
        tr = pick_train(train, ds)
        er = pick_errors(test, preds)
        names = [os.path.basename(str(r["image"])) for r in tr] + [r["image"] for r in er]
        links = copy_images(names, ds, drive_root, out_dir)
        missing += [n for n, v in links.items() if not v]
        for r in tr:
            img = os.path.basename(str(r["image"]))
            data_rows.append({"dataset": ds, "question_id": str(r.get("question_id", "")),
                              "question_type": r.get("question_type", "unknown"), "anh": img,
                              "_link": links.get(img), "question": r["question"], "answers": " || ".join(_answers(r))})
        for r in er:
            err_rows.append({**r, "dataset": ds, "anh": r["image"], "_link": links.get(r["image"])})
    for i, r in enumerate(data_rows, 1):
        r["stt"] = i
    for i, r in enumerate(err_rows, 1):
        r["stt"] = i
    xlsx = os.path.join(out_dir, "W07_audit_400.xlsx")
    write_xlsx(xlsx, data_rows, err_rows)
    rep = {"xlsx": xlsx.replace("\\", "/"), "data_audit_rows": len(data_rows), "error_audit_rows": len(err_rows),
           "images_missing": missing, "seed": SEED, "runs": RUNS}
    json.dump(rep, open(os.path.join(out_dir, "audit_manifest.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    return rep


def _selftest() -> None:
    import tempfile
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as td:
        drive, exp, out = os.path.join(td, "drive"), os.path.join(td, "exp"), os.path.join(td, "out")
        os.makedirs(os.path.join(drive, DATA["vivqa"]["images"]))
        os.makedirs(os.path.join(drive, "data/vitextvqa_official"))
        types = ["color", "location", "number", "object"]
        vtrain = [{"question_id": f"v{i}", "question": f"q{i}?", "answers": [f"a{i}"], "image": f"v{i % 50}.jpg",
                   "question_type": types[i % 4]} for i in range(200)]
        vtest = [{"question_id": f"vt{i}", "question": f"t{i}?", "answers": ["đỏ"], "image": f"v{i % 50}.jpg",
                  "question_type": types[i % 4]} for i in range(160)]
        for i in range(50):
            open(os.path.join(drive, DATA["vivqa"]["images"], f"v{i}.jpg"), "wb").write(b"img")
        ttrain = [{"question_id": f"x{i}", "question": f"chữ {i}?", "answers": [f"b{i}"], "image": f"x{i % 30}.jpg"} for i in range(150)]
        ttest = [{"question_id": f"xt{i}", "question": f"đọc {i}?", "answers": ["vinmart"], "image": f"x{i % 30}.jpg"} for i in range(160)]
        with zipfile.ZipFile(os.path.join(drive, DATA["vitextvqa"]["zip"]), "w") as zf:
            for i in range(29):                     # x29.jpg cố ý thiếu -> phải báo trong images_missing
                zf.writestr(f"ViTextVQA_images/x{i}.jpg", b"img")
        for ds, tr, te in (("vivqa", vtrain, vtest), ("vitextvqa", ttrain, ttest)):
            json.dump(tr, open(os.path.join(drive, DATA[ds]["train"]), "w", encoding="utf-8"), ensure_ascii=False)
            json.dump(te, open(os.path.join(drive, DATA[ds]["test"]), "w", encoding="utf-8"), ensure_ascii=False)
            for k, r in RUNS[ds].items():
                os.makedirs(os.path.join(exp, r))
                preds = [{"question_id": t["question_id"], "question": t["question"], "answers": t["answers"],
                          "prediction": t["answers"][0] if (k == "B2" and i % 3 == 0) else "sai"} for i, t in enumerate(te)]
                json.dump(preds, open(os.path.join(exp, r, "predictions.json"), "w", encoding="utf-8"), ensure_ascii=False)
        rep = build(drive, out, exp)
        assert rep["data_audit_rows"] == 200 and rep["error_audit_rows"] == 200, rep
        assert rep["images_missing"] == ["x29.jpg"], rep["images_missing"]
        wb = load_workbook(rep["xlsx"])
        assert wb.sheetnames == ["huong_dan", "data_audit", "error_audit"]
        da, ea = wb["data_audit"], wb["error_audit"]
        qt = [da.cell(i, 4).value for i in range(2, 102)]
        assert all(qt.count(t) == 25 for t in types), "ViVQA phải 25 mẫu/loại"
        hdr = [c.value for c in ea[1]]
        assert hdr[-5:] == ["loai_loi", "b2_dung_nghia", "hallucination", "ghi_chu", "nguoi_duyet"]
        b2 = hdr.index("B2") + 1
        assert all(ea.cell(i, b2).value == "sai" for i in range(2, 202)), "error_audit chỉ chứa câu B2 sai"
        assert ea.cell(2, hdr.index("anh") + 1).hyperlink is not None
        assert os.path.exists(os.path.join(out, "images", "vitextvqa", "x0.jpg"))
        assert len(ea.data_validations.dataValidation) == 3
        rep2 = build(drive, out, exp)                # chạy lại: tất định, không lỗi khi ảnh đã có
        assert rep2["data_audit_rows"] == 200
        wb2 = load_workbook(rep2["xlsx"])
        assert [c.value for c in wb2["error_audit"]["C"]] == [c.value for c in ea["C"]], "phải tất định"
    print("SELFTEST OK")


def main() -> None:
    ap = argparse.ArgumentParser(description="Tạo bảng duyệt tay W07 (data_audit + error_audit) kèm ảnh.")
    ap.add_argument("--drive_root", default="G:/My Drive/ViVQA-VLM")
    ap.add_argument("--out_dir", default=None, help="Mặc định <drive_root>/audit/W07")
    ap.add_argument("--exp_root", default=os.path.join(REPO_ROOT, "experiments"), help="Thư mục experiments có predictions")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        _selftest()
        return
    rep = build(a.drive_root, a.out_dir or os.path.join(a.drive_root, "audit", "W07"), a.exp_root)
    print(json.dumps({k: v for k, v in rep.items() if k != "runs"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
