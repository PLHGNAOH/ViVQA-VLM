"""
audit_w12.py — Audit tay lỗi của HAI MÔ HÌNH CHÍNH THỨC (W12, bước 4b): 100 lỗi B3-full + 50 lỗi C3.

Vấn đề: phân tích tự động (src/eval/error_analysis.py) chỉ so chuỗi. Nó không biết một câu "sai EM" thật ra có
đúng ý không, có phải gold sai không, hay model nói thứ không có trong ảnh (hallucination). Chỉ người xem ảnh mới
trả lời được. Luật ghi trước: docs/W12/README.md, mục "Error analysis" (đoạn "Audit tay").

  --make_sheet  chọn mẫu tất định (seed 42) trong các câu sai EM, chép ảnh, tạo Excel có dropdown.
                Phiếu MÙ: chỉ có câu hỏi, gold, câu trả lời của mô hình được audit. Không có kiểu lỗi / nguồn lỗi tự động,
                không có dự đoán của run khác, không có văn bản OCR.
  --score       đọc phiếu đã gán, tính % theo nhãn (Wilson CI95), % đúng ý, % hallucination, "EM theo người đọc";
                đối chiếu nhãn tay với nguồn lỗi tự động; so (chỉ mô tả) với audit W07 của B2.
  --selftest    kiểm thử offline.

Bộ nhãn = đúng bộ nhãn W07 (src/data/make_audit_sheet.py: ERROR_LABELS) -> so sánh được với audit B2.

Chạy (local, Drive for desktop):
    python -m src.eval.audit_w12 --make_sheet --drive_root "G:/My Drive/ViVQA-VLM" --out_dir "G:/My Drive/ViVQA-VLM/audit/W12"
    python -m src.eval.audit_w12 --score docs/W12/audit/W12_audit_150_labeled.xlsx --drive_root "G:/My Drive/ViVQA-VLM" \
        --out experiments/W12_error_audit
"""
from __future__ import annotations

import argparse
import json
import math
import os
import random
import sys
from collections import Counter
from typing import Any, Dict, List, Optional, Tuple

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from src.data.make_audit_sheet import ERROR_LABELS, copy_images  # noqa: E402
from src.eval.metrics import exact_match  # noqa: E402

SEED = 42
N = {"vitextvqa": 100, "vivqa": 50}
RUNS = {"vitextvqa": "W10_B3full_eval_qlora_ocr_vietocr_v1_r16_full_ep1_vitextvqa_testfull_bf16_L4_seed42",
        "vivqa": "W09_C3_eval_qlora_r16_vivqa_test_bf16_L4_seed42"}
MODEL = {"vitextvqa": "B3-full", "vivqa": "C3"}
TEST = {"vitextvqa": "data/vitextvqa_official/test.json", "vivqa": "data/vivqa/test.json"}
SHEET_NAME = "W12_audit_150.xlsx"
COLS = ["stt", "dataset", "question_id", "question_type", "anh", "question", "gold", "du_doan"]
HUMAN = ["loai_loi", "dung_y", "hallucination", "ghi_chu"]
YES_NO = ["Có", "Không"]
W07_FINAL = os.path.join(REPO_ROOT, "docs", "W07", "audit", "W07_error_audit_final.json")
GUIDE = [
    ("Mục đích", "Gán nhãn tay 150 câu mà mô hình CHÍNH THỨC trả lời sai EM: 100 câu B3-full (ViTextVQA full test) + "
                 "50 câu C3 (ViVQA test). Kết quả cho R6 §5 (error analysis) và §6 (hallucination)."),
    ("Cách làm", "Mở ảnh bằng link ở cột 'anh' (hoặc tìm theo tên file trong thư mục images/). So câu trả lời ở cột "
                 "'du_doan' với ẢNH và câu hỏi. Chọn nhãn ở 3 cột màu vàng bằng dropdown; ghi chú khi cần."),
    ("loai_loi", "Nguyên nhân CHÍNH (1 nhãn), cùng bộ nhãn với audit W07: " + " | ".join(ERROR_LABELS)),
    ("dung_y", "Có = người đọc chấp nhận câu trả lời là đúng ý (dù EM chấm sai). Không = sai thật."),
    ("hallucination", "Có = câu trả lời nói một thứ KHÔNG có trong ảnh (vật / chữ / số không hề xuất hiện). "
                      "Đọc nhầm một chữ có thật trong ảnh thì KHÔNG phải hallucination."),
    ("Quy tắc", "Gán MÙ: không mở file kết quả error analysis tự động, không xem dự đoán của mô hình khác trước khi gán "
                "xong. Không sửa cột không tô vàng. Dòng khó: chọn 'Khác' + ghi chú, không đoán. ≈ 40 giây / dòng."),
    ("Chọn mẫu", f"Tất định, seed {SEED}: ngẫu nhiên trong các câu sai EM (xếp question_id rồi random.sample)."),
]


def _load(path: str) -> Any:
    with open(path, encoding="utf-8") as fp:
        return json.load(fp)


def _records(path: str) -> List[Dict[str, Any]]:
    raw = _load(path)
    if isinstance(raw, dict):
        return raw["data"] if isinstance(raw.get("data"), list) else list(raw.values())
    return raw


def wilson(k: int, n: int, z: float = 1.959964) -> List[float]:
    if n == 0:
        return [0.0, 0.0]
    p = k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return [round(100 * (c - h) / d, 1), round(100 * (c + h) / d, 1)]


def pick(preds: List[Dict[str, Any]], n: int, seed: int = SEED) -> List[str]:
    wrong = sorted(str(p["question_id"]) for p in preds if not exact_match(p["prediction"], p["answers"]))
    return random.Random(seed).sample(wrong, min(n, len(wrong)))


def build_rows(preds: Dict[str, List[Dict[str, Any]]], tests: Dict[str, List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    rows = []
    for ds in ("vitextvqa", "vivqa"):
        by_p = {str(p["question_id"]): p for p in preds[ds]}
        by_t = {str(t["question_id"]): t for t in tests[ds]}
        for q in pick(preds[ds], N[ds]):
            p, t = by_p[q], by_t[q]
            rows.append({"dataset": ds, "question_id": q, "question_type": p.get("question_type", "unknown"),
                         "anh": os.path.basename(str(t["image"])), "question": p["question"],
                         "gold": " || ".join(p["answers"]), "du_doan": p["prediction"]})
    for i, r in enumerate(rows, 1):
        r["stt"] = i
    return rows


def write_sheet(path: str, rows: List[Dict[str, Any]], links: Dict[Tuple[str, str], str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    g = wb.active
    g.title = "huong_dan"
    for i, (k, v) in enumerate(GUIDE, 1):
        g.cell(i, 1, k).font = Font(bold=True)
        g.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    g.column_dimensions["A"].width, g.column_dimensions["B"].width = 16, 120
    ws = wb.create_sheet("error_audit")
    ws.append(COLS + HUMAN)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c, "") for c in COLS] + [""] * len(HUMAN))
        link = links.get((r["dataset"], r["anh"]))
        if link:
            cell = ws.cell(ws.max_row, COLS.index("anh") + 1)
            cell.hyperlink, cell.style = link, "Hyperlink"
    n = ws.max_row
    yellow = PatternFill("solid", fgColor="FFF2CC")
    lists = {"loai_loi": [x.replace(",", ";") for x in ERROR_LABELS], "dung_y": YES_NO, "hallucination": YES_NO}
    for j, h in enumerate(HUMAN, len(COLS) + 1):
        col = ws.cell(1, j).column_letter
        for i in range(2, n + 1):
            ws.cell(i, j).fill = yellow
        if h in lists:
            dv = DataValidation(type="list", formula1='"' + ",".join(lists[h]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{n}")
    for j, c in enumerate(COLS + HUMAN, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = 42 if c in ("question", "gold", "du_doan") else 15
    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def make_sheet(drive_root: str, out_dir: str, exp_root: str) -> Dict[str, Any]:
    preds = {ds: _load(os.path.join(exp_root, r, "predictions.json")) for ds, r in RUNS.items()}
    tests = {ds: _records(os.path.join(drive_root, p)) for ds, p in TEST.items()}
    rows = build_rows(preds, tests)
    links, missing = {}, []
    for ds in ("vitextvqa", "vivqa"):
        got = copy_images([r["anh"] for r in rows if r["dataset"] == ds], ds, drive_root, out_dir)
        missing += [n for n, v in got.items() if not v]
        links.update({(ds, n): v for n, v in got.items() if v})
    path = os.path.join(out_dir, SHEET_NAME)
    write_sheet(path, rows, links)
    man = {"xlsx": path.replace("\\", "/"), "rows": len(rows), "per_dataset": dict(Counter(r["dataset"] for r in rows)),
           "question_ids": [r["question_id"] for r in rows], "images_missing": missing, "seed": SEED, "runs": RUNS}
    with open(os.path.join(out_dir, "audit_manifest.json"), "w", encoding="utf-8") as fp:
        json.dump(man, fp, ensure_ascii=False, indent=2)
    return man


# =============================================================================
# Chấm phiếu
# =============================================================================
def read_sheet(path: str) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["error_audit"]
    head = [c.value for c in ws[1]]
    if head[:len(COLS) + len(HUMAN)] != COLS + HUMAN:
        raise ValueError(f"Cột phiếu khác khuôn: {head}")
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        d = dict(zip(head, r))
        d["question_id"] = str(d["question_id"])
        out.append(d)
    return out


def check_labels(rows: List[Dict[str, Any]], expect_ids: List[str]) -> List[str]:
    errs = []
    if [r["question_id"] for r in rows] != expect_ids:
        errs.append("Danh sách / thứ tự question_id khác phiếu gốc (seed 42)")
    labs = {x.replace(",", ";") for x in ERROR_LABELS}
    for r in rows:
        if r["loai_loi"] not in labs:
            errs.append(f"stt {r['stt']}: loai_loi trống hoặc lạ ({r['loai_loi']!r})")
        for k in ("dung_y", "hallucination"):
            if r[k] not in YES_NO:
                errs.append(f"stt {r['stt']}: {k} trống hoặc lạ ({r[k]!r})")
    return errs


def summarize(rows: List[Dict[str, Any]], em_pct: float) -> Dict[str, Any]:
    n = len(rows)
    lab = Counter(r["loai_loi"] for r in rows)
    ok = sum(r["dung_y"] == "Có" for r in rows)
    hal = sum(r["hallucination"] == "Có" for r in rows)
    err_rate = 100 - em_pct
    ok_ci = wilson(ok, n)
    return {"n": n,
            "error_type_pct": {k: round(100 * v / n, 1) for k, v in lab.most_common()},
            "error_type_ci95": {k: wilson(v, n) for k, v in lab.most_common()},
            "correct_in_meaning_pct": round(100 * ok / n, 1), "correct_in_meaning_ci95": ok_ci,
            "hallucination_pct": round(100 * hal / n, 1), "hallucination_ci95": wilson(hal, n),
            "em_pct": round(em_pct, 2),
            "reader_accuracy_pct": round(em_pct + err_rate * ok / n / 1, 2) if n else None,
            "reader_accuracy_ci95": [round(em_pct + err_rate * x / 100, 2) for x in ok_ci],
            "hallucination_of_all_answers_pct": round(err_rate * hal / n, 2) if n else None}


def score(rows: List[Dict[str, Any]], preds: Dict[str, List[Dict[str, Any]]],
          auto: Optional[Dict[str, Any]] = None, w07: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    out: Dict[str, Any] = {"by_dataset": {}}
    for ds in ("vitextvqa", "vivqa"):
        rr = [r for r in rows if r["dataset"] == ds]
        em = 100 * sum(exact_match(p["prediction"], p["answers"]) for p in preds[ds]) / len(preds[ds])
        out["by_dataset"][ds] = {"model": MODEL[ds], **summarize(rr, em)}
        if w07:
            w = w07["stats"].get(ds, {})
            out["by_dataset"][ds]["w07_B2"] = {k: w.get(k) for k in ("n", "correct_in_meaning_pct", "hallucination_pct",
                                                                     "error_type_pct")}
    if auto:
        src = {}
        for ds_key, ex_key in (("vitext", "vitextvqa"),):
            for r in rows:
                if r["dataset"] == ex_key and r["question_id"] in auto:
                    src.setdefault(auto[r["question_id"]], Counter())[r["loai_loi"]] += 1
        out["vitextvqa_human_label_x_auto_source"] = {k: dict(v) for k, v in src.items()}
    return out


def auto_sources(drive_root: str, exp_root: str) -> Dict[str, str]:
    """Nguồn lỗi tự động của B3-full cho từng câu (tính lại bằng error_analysis, cùng định nghĩa)."""
    from src.data.ocr import load_ocr_cache
    from src.eval.error_analysis import error_source, is_right, ocr_status, ocr_text_of, pred_in_ocr
    preds = _load(os.path.join(exp_root, RUNS["vitextvqa"], "predictions.json"))
    test = {str(t["question_id"]): t for t in _records(os.path.join(drive_root, TEST["vitextvqa"]))}
    cache = load_ocr_cache(os.path.join(drive_root, "data/vitextvqa_official/ocr/paddledet_vietocr.json"))["images"]
    out = {}
    for p in preds:
        if is_right(p):
            continue
        txt = ocr_text_of(cache[os.path.basename(str(test[str(p["question_id"])]["image"]))])
        out[str(p["question_id"])] = error_source(ocr_status(txt, p["answers"]), pred_in_ocr(txt, p["prediction"]))
    return out


def report(s: Dict[str, Any]) -> str:
    L = ["# Audit tay lỗi mô hình chính thức (W12, bước 4b)", "",
         "| | B3-full (ViTextVQA) | C3 (ViVQA) | B2 W07 ViTextVQA | B2 W07 ViVQA |", "|---|---|---|---|---|"]
    vt, vv = s["by_dataset"]["vitextvqa"], s["by_dataset"]["vivqa"]
    w = lambda d, k: (d.get("w07_B2") or {}).get(k, "—")
    L.append(f"| n lỗi duyệt | {vt['n']} | {vv['n']} | {w(vt, 'n')} | {w(vv, 'n')} |")
    L.append(f"| % đúng ý [CI95] | {vt['correct_in_meaning_pct']} {vt['correct_in_meaning_ci95']} | "
             f"{vv['correct_in_meaning_pct']} {vv['correct_in_meaning_ci95']} | {w(vt, 'correct_in_meaning_pct')} | {w(vv, 'correct_in_meaning_pct')} |")
    L.append(f"| % hallucination [CI95] | {vt['hallucination_pct']} {vt['hallucination_ci95']} | "
             f"{vv['hallucination_pct']} {vv['hallucination_ci95']} | {w(vt, 'hallucination_pct')} | {w(vv, 'hallucination_pct')} |")
    L.append(f"| EM → EM theo người đọc [CI95] | {vt['em_pct']} → {vt['reader_accuracy_pct']} {vt['reader_accuracy_ci95']} | "
             f"{vv['em_pct']} → {vv['reader_accuracy_pct']} {vv['reader_accuracy_ci95']} | | |")
    L.append(f"| Hallucination / mọi câu trả lời | {vt['hallucination_of_all_answers_pct']}% | {vv['hallucination_of_all_answers_pct']}% | | |")
    for ds, d in (("ViTextVQA · B3-full", vt), ("ViVQA · C3", vv)):
        L += ["", f"## Kiểu lỗi — {ds}", ""] + [f"- {k}: {v}% {d['error_type_ci95'][k]}" for k, v in d["error_type_pct"].items()]
    if s.get("vitextvqa_human_label_x_auto_source"):
        L += ["", "## Nhãn tay × nguồn lỗi tự động (B3-full)", ""]
        for k, v in s["vitextvqa_human_label_x_auto_source"].items():
            L.append(f"- **{k}**: " + "; ".join(f"{a} {b}" for a, b in sorted(v.items(), key=lambda x: -x[1])))
    return "\n".join(L) + "\n"


# =============================================================================
def _selftest() -> None:
    import tempfile
    assert wilson(0, 10)[0] == 0.0 and wilson(5, 10) == [23.7, 76.3]
    mk = lambda ds, k, right: [{"question_id": f"{ds}{i}", "question": f"q{i}", "answers": ["a"],
                                "prediction": "a" if i < right else "b", "question_type": "unknown"} for i in range(k)]
    preds = {"vitextvqa": mk("t", 300, 150), "vivqa": mk("v", 200, 120)}
    tests = {ds: [{"question_id": p["question_id"], "image": f"{p['question_id']}.jpg"} for p in v] for ds, v in preds.items()}
    rows = build_rows(preds, tests)
    assert len(rows) == 150 and [r["stt"] for r in rows] == list(range(1, 151))
    assert all(r["du_doan"] == "b" for r in rows), "chỉ chọn câu sai EM"
    assert rows == build_rows(preds, tests), "tất định"
    assert set(COLS + HUMAN) >= {"du_doan"} and "B2" not in COLS and "source" not in COLS, "phiếu phải mù"
    with tempfile.TemporaryDirectory() as td:
        p = os.path.join(td, "s.xlsx")
        write_sheet(p, rows, {})
        back = read_sheet(p)
        assert [r["question_id"] for r in back] == [r["question_id"] for r in rows]
        assert check_labels(back, [r["question_id"] for r in rows]), "phiếu trống phải báo lỗi"
        from openpyxl import load_workbook
        wb = load_workbook(p)
        ws = wb["error_audit"]
        for i in range(2, ws.max_row + 1):
            ws.cell(i, len(COLS) + 1, ERROR_LABELS[0] if i % 2 else ERROR_LABELS[4])
            ws.cell(i, len(COLS) + 2, "Có" if i % 2 else "Không")
            ws.cell(i, len(COLS) + 3, "Không" if i % 2 else "Có")
        wb.save(p)
        back = read_sheet(p)
        assert not check_labels(back, [r["question_id"] for r in rows])
        auto = {r["question_id"]: "ocr_upstream" for r in back if r["dataset"] == "vitextvqa"}
        s = score(back, preds, auto=auto, w07=_load(W07_FINAL) if os.path.exists(W07_FINAL) else None)
        vt = s["by_dataset"]["vitextvqa"]
        assert vt["n"] == 100 and vt["em_pct"] == 50.0
        assert abs(vt["correct_in_meaning_pct"] + vt["hallucination_pct"] - 100) < 1e-9
        assert vt["reader_accuracy_pct"] == round(50 + 50 * vt["correct_in_meaning_pct"] / 100, 2)
        assert sum(s["vitextvqa_human_label_x_auto_source"]["ocr_upstream"].values()) == 100
        assert "EM theo người đọc" in report(s)
    print("audit_w12 selftest OK")


def main() -> int:
    ap = argparse.ArgumentParser(description="Audit tay lỗi mô hình chính thức (W12, bước 4b)")
    ap.add_argument("--make_sheet", action="store_true")
    ap.add_argument("--score", metavar="LABELED_XLSX")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--drive_root")
    ap.add_argument("--out_dir")
    ap.add_argument("--exp_root", default=os.path.join(REPO_ROOT, "experiments"))
    ap.add_argument("--out")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if a.selftest:
        _selftest()
        return 0
    if a.make_sheet:
        if not (a.drive_root and a.out_dir):
            ap.error("--make_sheet cần --drive_root --out_dir")
        if os.path.exists(os.path.join(a.out_dir, SHEET_NAME)):
            print(f"⏭ {SHEET_NAME} đã có ở {a.out_dir} -> KHÔNG tạo lại (tránh ghi đè nhãn đã gán)")
            return 0
        m = make_sheet(a.drive_root, a.out_dir, a.exp_root)
        print(f"✔ {m['rows']} dòng {m['per_dataset']} | ảnh thiếu: {len(m['images_missing'])} -> {m['xlsx']}")
        return 0
    if a.score:
        if not (a.drive_root and a.out):
            ap.error("--score cần --drive_root --out")
        rows = read_sheet(a.score)
        preds = {ds: _load(os.path.join(a.exp_root, r, "predictions.json")) for ds, r in RUNS.items()}
        tests = {ds: _records(os.path.join(a.drive_root, p)) for ds, p in TEST.items()}
        errs = check_labels(rows, [r["question_id"] for r in build_rows(preds, tests)])
        if errs:
            print("❌ Phiếu chưa hợp lệ:\n" + "\n".join(errs[:30]))
            return 1
        s = score(rows, preds, auto=auto_sources(a.drive_root, a.exp_root),
                  w07=_load(W07_FINAL) if os.path.exists(W07_FINAL) else None)
        os.makedirs(a.out, exist_ok=True)
        with open(os.path.join(a.out, "audit_summary.json"), "w", encoding="utf-8", newline="\n") as fp:
            json.dump(s, fp, ensure_ascii=False, indent=1)
        md = report(s)
        with open(os.path.join(a.out, "audit_summary.md"), "w", encoding="utf-8", newline="\n") as fp:
            fp.write(md)
        print(md)
        return 0
    ap.error("chọn --make_sheet, --score hoặc --selftest")
    return 2


if __name__ == "__main__":
    sys.exit(main())
