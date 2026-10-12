"""
compute_error_audit.py — Tính số liệu error audit W07 (200 lỗi EM của B2) từ file gán nhãn của nhóm.

Nguồn nhãn (docs/W07/audit/labels/):
  W07_audit_nguoi2/3/4.xlsx   nhãn người độc lập (Huân, Trường, Toàn). 20 dòng 'hieu_chinh' là bộ chung.
  W07_audit_nguoi1_Hoang.xlsx 42 dòng 'rieng' = AI pre-annotated (Claude v1) → Huân → Hoàng duyệt (AI-assisted).
  ai_prelabels_v2.json        nhãn gợi ý của AI (v2) + 10 dòng Hoàng tự gán trước khi thấy nhãn AI (stt 1–10).
  W07_adjudication.xlsx       Hoàng phân xử các dòng người≠AI ở b2_dung_nghia/hallucination + dòng 2 người lệch.
Quy tắc:
  - Nhãn chính = nhãn người; 20 dòng chung: Huân & Trường (Toàn bị loại ở bộ chung vì ghi chú trùng nguyên văn Trường).
  - Dòng có trong file phân xử: dùng lựa chọn của người phân xử (Người / AI / Khác + final_*).
  - κ giữa người: Cohen, Huân vs Trường trên 20 dòng chung. κ người–AI: các dòng có nhãn người độc lập.
  - Sensitivity: bỏ 42 dòng AI-assisted.
Chạy: python docs/W07/audit/compute_error_audit.py [--out docs/W07/audit/W07_error_audit_final.json]
"""
import argparse, json, math, os
from collections import Counter
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); LAB = os.path.join(HERE, 'labels')
F = ['loai_loi', 'b2_dung_nghia', 'hallucination']

def load(name):
    ws = openpyxl.load_workbook(os.path.join(LAB, name), data_only=True)['error_audit']
    H = [c.value for c in ws[1]]
    return {int(r[0]): dict(zip(H, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}

def cohen(a, b):
    n = len(a); po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b); pe = sum(ca[k] * cb[k] for k in set(a) | set(b)) / n / n
    return po, ((po - pe) / (1 - pe) if pe < 1 else float('nan'))

def wilson(k, n, z=1.959964):
    p = k / n; d = 1 + z * z / n; c = p + z * z / (2 * n); h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return [round(100 * (c - h) / d, 1), round(100 * (c + h) / d, 1)]

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--out', default=os.path.join(HERE, 'W07_error_audit_final.json')); a = ap.parse_args()
    ai = json.load(open(os.path.join(LAB, 'ai_prelabels_v2.json'), encoding='utf-8'))
    NAME = ai['code_names']; INV = {v: k for k, v in NAME.items()}
    AI = {int(s): (NAME[d['code']], d['b2_dung_nghia'], d['hallucination']) for s, d in ai['labels'].items()}
    H0 = {int(s): tuple(d[f] for f in F) for s, d in ai['hoang_rows_1_10'].items()}
    T = {k: load(f'W07_audit_{k}.xlsx') for k in ['nguoi1_Hoang', 'nguoi2', 'nguoi3', 'nguoi4']}
    lab = lambda d: tuple(d[f] for f in F)
    cal = sorted(s for s, d in T['nguoi2'].items() if d['phan'] == 'hieu_chinh')
    human, src = {}, {}
    for s in range(1, 201):
        if s <= 10: human[s], src[s] = H0[s], 'Hoàng (independent)'
        elif s in cal:
            x, y = lab(T['nguoi2'][s]), lab(T['nguoi3'][s])
            human[s], src[s] = (x if x == y else None), 'Huân+Trường (independent, shared set)'
        else:
            for k, who in [('nguoi2', 'Huân'), ('nguoi3', 'Trường'), ('nguoi4', 'Toàn'), ('nguoi1_Hoang', 'AI-assisted (Claude v1 → Huân → Hoàng)')]:
                d = T[k].get(s)
                if d and d['phan'] == 'rieng':
                    h = lab(d); human[s], src[s] = (h if all(h) else None), who
    ws = openpyxl.load_workbook(os.path.join(LAB, 'W07_adjudication.xlsx'), data_only=True)['adjudication']
    Hd = [c.value for c in ws[1]]; ADJ = {int(r[0]): dict(zip(Hd, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    final = {}
    for s in range(1, 201):
        if s in ADJ:
            c = ADJ[s]['chon']
            final[s] = AI[s] if c == 'AI' else (human[s] if c == 'Người' else (ADJ[s]['final_loai_loi'], ADJ[s]['final_b2_dung_nghia'], ADJ[s]['final_hallucination']))
        else:
            final[s] = human[s]
    assert all(final[s] and all(final[s]) for s in final), 'thiếu nhãn cuối'
    def summ(ss):
        n = len(ss); c = Counter(INV[final[s][0]] for s in ss); k_d = sum(final[s][1] == 'Có' for s in ss); k_h = sum(final[s][2] == 'Có' for s in ss)
        return dict(n=n, correct_in_meaning_pct=round(100 * k_d / n, 1), correct_in_meaning_ci95=wilson(k_d, n),
                    hallucination_pct=round(100 * k_h / n, 1), hallucination_ci95=wilson(k_h, n),
                    error_type_pct={k: round(100 * v / n, 1) for k, v in sorted(c.items(), key=lambda x: -x[1])})
    S = {'all': summ(list(range(1, 201))), 'vivqa': summ(list(range(1, 101))), 'vitextvqa': summ(list(range(101, 201))),
         'all_without_ai_assisted': summ([s for s in range(1, 201) if not src[s].startswith('AI')])}
    S['iaa_shared_set_Huan_vs_Truong'] = {f: dict(zip(['agree', 'kappa'], [round(x, 3) for x in cohen([T['nguoi2'][s][f] for s in cal], [T['nguoi3'][s][f] for s in cal])])) for f in F[:2]}
    ind = [s for s in range(1, 201) if human[s] and not src[s].startswith('AI')]
    S['human_vs_ai'] = {'n': len(ind), **{f: dict(zip(['agree', 'kappa'], [round(x, 3) for x in cohen([human[s][i] for s in ind], [AI[s][i] for s in ind])])) for i, f in enumerate(F)}}
    S['adjudication'] = {'n': len(ADJ), 'choice': dict(Counter(d['chon'] for d in ADJ.values()))}
    S['provenance'] = dict(Counter(src.values()))
    json.dump({'stats': S, 'final_labels': {s: list(v) for s, v in final.items()}, 'source': src}, open(a.out, 'w', encoding='utf-8', newline='\n'), ensure_ascii=False, indent=1)
    for k, v in S.items(): print(k, v)
    print('Ghi ->', a.out)

if __name__ == '__main__':
    main()
